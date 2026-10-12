"""Editable Word export of the saved ONYX report."""
from io import BytesIO
from pathlib import Path


def build_docx(mission, report, entries, upload_directory):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(1.1)
    section.bottom_margin = Inches(0.8)
    section.header_distance = Inches(0.3)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.space_after = Pt(6)
    document.styles['Title'].font.size = Pt(24)
    from docx.oxml.ns import qn
    for style_name in ('Title', 'Normal'):
        properties = document.styles[style_name].element.find(qn('w:pPr'))
        if properties is not None:
            for border in list(properties.findall(qn('w:pBdr'))):
                properties.remove(border)
    logo_candidates = [Path(__file__).parent / 'static/onyx-logo.png', Path(__file__).parent.parent / 'web/onyx-logo.png']
    logo = next((path for path in logo_candidates if path.is_file()), None)
    if logo:
        section.header.paragraphs[0].add_run().add_picture(str(logo), width=Inches(1.2))
    else:
        section.header.paragraphs[0].text = 'ONYX'
    document.add_paragraph('Rapport de fin de journée', 'Title')
    document.add_paragraph(mission['title'])
    document.add_paragraph(f"Client : {mission['client_name']}\nRéférence : {mission['reference']}\nDate de mission : {mission['mission_date']}")
    for line in report['content'].splitlines():
        line = line.strip()
        if not line or line.startswith('ONYX — RAPPORT'):
            continue
        if line.startswith(('Mission :', 'Client :', 'Référence :', 'Date :')):
            continue
        if line.isupper() and len(line) < 100:
            document.add_paragraph(line.capitalize(), 'Heading 1')
        elif line.startswith('• '):
            document.add_paragraph(line[2:], 'List Bullet')
        else:
            document.add_paragraph(line)
    photos = [entry for entry in entries if entry['entry_type'] == 'photo']
    if photos:
        document.add_page_break()
        document.add_paragraph('Annexes photographiques', 'Heading 1')
        root = Path(upload_directory).resolve()
        for index, entry in enumerate(photos, 1):
            path = (root / entry['file_path']).resolve()
            caption = document.add_paragraph(f"Photo {index} — {entry['original_name']}")
            caption.paragraph_format.keep_with_next = True
            if root not in path.parents or not path.is_file():
                document.add_paragraph('Pièce jointe indisponible.')
                continue
            try:
                from PIL import Image
                with Image.open(path) as source:
                    source.thumbnail((1400, 1000))
                    buffer = BytesIO()
                    source.convert('RGB').save(buffer, format='JPEG')
                    buffer.seek(0)
                    document.add_picture(buffer, width=Inches(6))
            except (OSError, ValueError):
                document.add_paragraph('Consulter la photo originale dans la mission.')
    section.footer.paragraphs[0].text = f"ONYX · {mission['reference']} · Rapport modifiable"
    output = BytesIO()
    document.save(output)
    return output.getvalue()
